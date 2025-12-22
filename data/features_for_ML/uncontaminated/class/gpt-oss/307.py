import os
import re
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Tuple

import requests
from tenacity import retry, stop_after_attempt, wait_fixed

# Simple retry decorator that retries on specified exceptions
def retry_on_exceptions(exceptions):
    return retry(
        reraise=True,
        stop=stop_after_attempt(3),
        wait=wait_fixed(1),
        retry=retry_if_exception_type(exceptions),
    )

# Helper for tenacity
def retry_if_exception_type(exceptions):
    def predicate(retry_state):
        return isinstance(retry_state.outcome.exception(), exceptions)
    return predicate


class DocumentProcessingToolkit:
    r"""A class representing a toolkit for processing document and return the content of the document.

    This class provides method for processing docx, pdf, pptx, etc. It cannot process excel files.
    """

    def __init__(self, cache_dir: Optional[str] = None):
        self.cache_dir = Path(cache_dir) if cache_dir else Path(tempfile.gettempdir()) / "doc_toolkit_cache"
        self.cache_dir.mkdir(parents=True, exist_ok=True)

    @retry_on_exceptions((requests.RequestException,))
    def extract_document_content(self, document_path: str) -> Tuple[bool, str]:
        """
        Extracts text content from a document file or URL.

        Returns:
            Tuple[bool, str]: (True, content) on success, (False, error_message) on failure.
        """
        try:
            # Determine if input is a URL
            if self._is_webpage(document_path):
                local_path = self._download_file(document_path)
            else:
                local_path = Path(document_path)

            if not local_path.exists():
                return False, f"File not found: {local_path}"

            ext = local_path.suffix.lower()
            content_parts: List[str] = []

            if ext == ".zip":
                extracted_files = self._unzip_file(str(local_path))
                for f in extracted_files:
                    ok, part = self.extract_document_content(f)
                    if ok:
                        content_parts.append(part)
                return True, "\n".join(content_parts)

            if ext == ".docx":
                try:
                    import docx
                except ImportError:
                    return False, "python-docx is required for .docx files."
                doc = docx.Document(str(local_path))
                for para in doc.paragraphs:
                    content_parts.append(para.text)
                return True, "\n".join(content_parts)

            if ext == ".pdf":
                try:
                    import PyPDF2
                except ImportError:
                    return False, "PyPDF2 is required for .pdf files."
                with open(str(local_path), "rb") as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        content_parts.append(page.extract_text() or "")
                return True, "\n".join(content_parts)

            if ext == ".pptx":
                try:
                    import pptx
                except ImportError:
                    return False, "python-pptx is required for .pptx files."
                prs = pptx.Presentation(str(local_path))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            content_parts.append(shape.text)
                return True, "\n".join(content_parts)

            if ext in {".xls", ".xlsx"}:
                return False, "Excel files are not supported."

            return False, f"Unsupported file type: {ext}"

        except Exception as e:
            return False, f"Error processing document: {str(e)}"

    def _is_webpage(self, url: str) -> bool:
        return bool(re.match(r"^https?://", url))

    def _download_file(self, url: str) -> Path:
        local_filename = self.cache_dir / Path(url).name
        with requests.get(url, stream=True, timeout=10) as r:
            r.raise_for_status()
            with open(local_filename, "wb") as f:
                shutil.copyfileobj(r.raw, f)
        return local_filename

    def _get_formatted_time(self) -> str:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def _unzip_file(self, zip_path: str) -> List[str]:
        extracted_paths: List[str] = []
        with zipfile.ZipFile(zip_path, "r") as zip_ref:
            extract_dir = self.cache_dir / Path(zip_path).stem
            extract_dir.mkdir(parents=True, exist_ok=True)
            zip_ref.extractall(extract_dir)
            for root, _, files in os.walk(extract_dir):
                for file in files:
                    extracted_paths.append(str(Path(root) / file))
        return extracted_paths